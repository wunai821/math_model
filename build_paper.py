from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE

ROOT = Path(r"D:\D题")
OUT = ROOT / "output"
OUT.mkdir(exist_ok=True)
FIG = ROOT / "answer" / "figures"
OUTPUT = OUT / "D题论文初稿及答辩问答.docx"

doc = Document()
sec = doc.sections[0]
sec.top_margin = Cm(2.4)
sec.bottom_margin = Cm(2.2)
sec.left_margin = Cm(2.6)
sec.right_margin = Cm(2.4)
sec.header_distance = Cm(1.2)
sec.footer_distance = Cm(1.2)

def set_font(run, east="宋体", latin="Times New Roman", size=10.5, bold=False, color="000000"):
    run.font.name = latin
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), latin)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), latin)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)

def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill)

def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, val in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")

def set_table_borders(table, color="D9D9D9", size="6"):
    tblPr = table._tbl.tblPr
    borders = tblPr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tblPr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:color"), color)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Times New Roman"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
normal.font.size = Pt(10.5)
normal.paragraph_format.line_spacing = 1.5
normal.paragraph_format.space_after = Pt(3)
normal.paragraph_format.first_line_indent = Cm(0.74)

title = styles["Title"]
title.font.name = "Times New Roman"
title._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
title.font.size = Pt(22)
title.font.bold = True
title.font.color.rgb = RGBColor(0, 0, 0)
title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_after = Pt(14)
# 清除Word内置Title样式可能携带的底边框
title_ppr = title._element.get_or_add_pPr()
for old in list(title_ppr.findall(qn("w:pBdr"))):
    title_ppr.remove(old)
title_borders = OxmlElement("w:pBdr")
for edge in ("top", "left", "bottom", "right", "between", "bar"):
    el = OxmlElement(f"w:{edge}")
    el.set(qn("w:val"), "nil")
    title_borders.append(el)
title_ppr.append(title_borders)

for name, size in [("Heading 1", 15), ("Heading 2", 13), ("Heading 3", 11.5)]:
    st = styles[name]
    st.font.name = "Times New Roman"
    st._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor(0, 0, 0)
    st.paragraph_format.space_before = Pt(10 if name == "Heading 1" else 7)
    st.paragraph_format.space_after = Pt(5)
    st.paragraph_format.keep_with_next = True

if "Caption CN" not in styles:
    cap = styles.add_style("Caption CN", WD_STYLE_TYPE.PARAGRAPH)
    cap.font.name = "Times New Roman"
    cap._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    cap.font.size = Pt(9)
    cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_before = Pt(4)
    cap.paragraph_format.space_after = Pt(8)
    cap.paragraph_format.keep_with_next = True

def body(text, indent=True, bold_lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.first_line_indent = Cm(0.74) if indent else Cm(0)
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_lead):])
        set_font(r2)
    else:
        r = p.add_run(text)
        set_font(r)
    return p

def formula(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_together = True
    r = p.add_run(text)
    set_font(r, east="Cambria Math", latin="Cambria Math", size=10.5)
    return p

def bullets(items, numbered=False):
    for item in items:
        p = doc.add_paragraph(style="List Number" if numbered else "List Bullet")
        p.paragraph_format.left_indent = Cm(0.74)
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(item)
        set_font(r)

def add_table(headers, rows, widths=None, font_size=9):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    hdr = table.rows[0].cells
    for j, h in enumerate(headers):
        hdr[j].text = str(h)
        set_cell_shading(hdr[j], "1F4E78")
        hdr[j].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(hdr[j])
        p = hdr[j].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
        for r in p.runs:
            set_font(r, east="黑体", size=font_size, bold=True, color="FFFFFF")
    for i, row in enumerate(rows):
        cells = table.add_row().cells
        for j, val in enumerate(row):
            cells[j].text = str(val)
            cells[j].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cells[j])
            if i % 2 == 1:
                set_cell_shading(cells[j], "F4F8FC")
            p = cells[j].paragraphs[0]
            p.paragraph_format.first_line_indent = Cm(0)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if len(str(val)) > 14 else WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                set_font(r, size=font_size)
        if widths:
            for j, w in enumerate(widths):
                cells[j].width = Cm(w)
    for row in table.rows:
        trPr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        cant_split.set(qn("w:val"), "1")
        trPr.append(cant_split)
    header_pr = table.rows[0]._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "1")
    header_pr.append(repeat)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table

def add_figure(filename, caption, width=15.5):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(5)
    run = p.add_run()
    run.add_picture(str(FIG / filename), width=Cm(width))
    c = doc.add_paragraph(caption, style="Caption CN")
    for r in c.runs:
        set_font(r, size=9)

def page_break():
    doc.add_page_break()

def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fldChar1 = OxmlElement("w:fldChar")
    fldChar1.set(qn("w:fldCharType"), "begin")
    instrText = OxmlElement("w:instrText")
    instrText.set(qn("xml:space"), "preserve")
    instrText.text = "PAGE"
    fldChar2 = OxmlElement("w:fldChar")
    fldChar2.set(qn("w:fldCharType"), "end")
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)
    set_font(run, size=9)

add_page_number(sec.footer.paragraphs[0])

# 摘要页
p = doc.add_paragraph(style="Title")
p.add_run("离散时频资源下用频计划冲突检测与多目标消解")

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(14)
r = p.add_run("D题论文初稿及答辩问答")
set_font(r, east="黑体", size=12, bold=True)

doc.add_heading("摘要", level=1)
body("针对周期性用频计划在离散时频资源上的冲突检测、参数消解与容量扩展问题，本文建立了从几何判定、冲突图分析到约束规划优化的统一建模框架。首先，将每项计划展开为一组左闭右开的周期时间区间，以频段交叠和任一周期时段交叠同时成立作为冲突的充要条件；遍历150项计划的11175个无序装备对，并用独立的单位时频网格法交叉验证。结果检测出297对冲突，涉及148台装备，其中B-C类冲突181对，占60.94%。")
body("其次，针对问题2构造相对时间平移安全差集，将周期时段相容性压缩为有限的相对位移关系，并采用CP-SAT进行七层字典序优化。模型以撤销总数、A类撤销数、B类撤销数、调整总数、A类调整数、B类调整数和归一化调整幅度为顺序逐层优化。所得方案保留18项、调整126项、撤销6项，最终执行144项计划且残余冲突为0；其中最少撤销6项以及A类零撤销已获最优性证明。")
body("在问题3中，固定问题2的无冲突计划，并把“不增加时频资源”解释为频域[0,100)、时域[0,643)。通过完整枚举C类计划的频率与首次时间起点，建立候选-资源0-1集合打包模型，求得最多新增142台C类装备，目标值与上界一致。问题4进一步允许C类间隔在±10个时间单位内调整，在原样、频移、时移、间隔调整和撤销候选之间进行互斥选择，得到保留11项、调整135项、撤销4项、执行146项的无冲突方案。该方案将撤销数由问题2的6项降至4项，但当前下界为2，故只报告为限时可行解。最后，独立验证器从原始数据重建全部重复时段，核查单参数限制、边界、类别权限、冲突及提交表格式，形成求解与验证相分离的结果闭环。")
body("本文方法兼具判定精确性、目标解释性与工程可复现性。研究表明，间隔调节能显著增加消解自由度；同时，资源窗口定义与最优性状态必须在结论中单独说明，不能以可行解替代最优解。")

p = doc.add_paragraph()
p.paragraph_format.first_line_indent = Cm(0)
r = p.add_run("关键词：")
set_font(r, east="黑体", bold=True)
r = p.add_run("时频冲突；周期区间；冲突图；约束规划；字典序优化；集合打包")
set_font(r)

add_figure("图11_四问建模与验证流程.png", "图1  四问建模与独立验证流程", width=15.7)

page_break()

doc.add_heading("1 问题重述", level=1)
body("系统提供100个离散频段，时间也按基本单位离散。每台装备提交频段区间、首次使用时间区间、相邻两次使用的空闲间隔和使用次数。若两台装备在某次使用中同时占用相交的时间区间和频段区间，则发生时频冲突。题目要求完成四项任务：检测全部冲突；在有限频移和时移范围内消解冲突；基于问题2结果评估还能新增的C类装备容量；允许C类间隔调整后重新消解冲突。")
bullets([
    "问题1：识别全部冲突装备对，统计冲突规模与类别结构。",
    "问题2：每台装备只能频移、时移、撤销或不变，且优先保护高等级计划，给出无冲突方案。",
    "问题3：固定问题2方案，在不增加资源且不限制新增计划平移幅度时，最大化新增C类装备数量。",
    "问题4：允许C类装备单独调整使用间隔，差异不超过10Δt，重新获得无冲突方案。"
])

add_figure("图1_原始时频资源占用.png", "图2  原始150项计划的时频资源占用", width=15.8)

doc.add_heading("2 问题分析", level=1)
doc.add_heading("2.1 冲突的本质", level=2)
body("冲突不是单独的频率重叠或时间重叠，而是二维时频矩形同时相交。由于每台装备会周期性重复使用，首次时间区间只是一个基元，必须展开全部使用次数后判断。将区间规定为左闭右开，可使相邻计划在端点处无重叠，并与整数网格占用保持一致。")

doc.add_heading("2.2 四问之间的递进关系", level=2)
body("问题1给出冲突图与冲突判据；问题2在固定动作集合内寻找无冲突布局；问题3把问题2结果视为不可移动的背景占用，对同质C类周期块做最大集合打包；问题4则回到原始150项计划，在问题2动作之外增加C类间隔变化。四问共享同一套周期展开、边界检查和独立冲突验证程序。")

doc.add_heading("2.3 难点与处理原则", level=2)
bullets([
    "周期展开会使单台计划对应多个矩形，需避免只检查首次时间区间。",
    "调整后可能产生原始冲突图之外的新冲突，因此无冲突约束必须覆盖所有装备对。",
    "题意给出多个定性目标而没有数值权重，采用字典序优化比任意加权和更易解释。",
    "问题3的资源时间上界未显式给出，必须明确解释并讨论结论的适用范围。",
    "限时求解返回的FEASIBLE只证明可行，只有目标值与界相等或状态为OPTIMAL时才能声称最优。"
])

doc.add_heading("3 模型假设", level=1)
bullets([
    "所有输入端点、间隔、使用次数及允许调整量均以Δf、Δt为单位取整数。",
    "频段与时间区间均采用左闭右开形式，端点相接不算冲突。",
    "空闲间隔指前一次使用结束到下一次使用开始之间的时长。",
    "问题2不额外设置统一结束时刻，只要求首次时间非负并满足题定平移幅度。",
    "问题3将不增加资源解释为频域[0,100)、时域[0,643)，其中643为原计划全部重复使用的最晚结束时刻。",
    "新增C类装备继承附件中C类统一需求：频宽3、单次时长2、间隔8、使用12次。",
    "问题4不把问题3的窗口假设外推为统一约束；若另行规定结束上限，应重新求解。"
])

page_break()
doc.add_heading("4 符号说明", level=1)
add_table(
    ["符号", "含义", "取值或单位"],
    [
        ("i,j", "装备索引", "1,…,150"),
        ("k,l", "周期使用次数索引", "非负整数"),
        ("[aᵢ,bᵢ)", "装备i的频段区间", "频率格"),
        ("[sᵢ,eᵢ)", "装备i首次时间区间", "时间格"),
        ("dᵢ", "单次使用时长 eᵢ-sᵢ", "正整数"),
        ("gᵢ,nᵢ", "空闲间隔、使用次数", "非负整数、正整数"),
        ("uᵢ,vᵢ,hᵢ", "频移、首次时间平移、间隔变化", "整数格"),
        ("zᵢ", "撤销指示变量", "0或1"),
        ("cᵢ", "是否调整", "0或1"),
        ("Sᵢⱼ", "装备对的安全相对时移集合", "{-10,…,10}的子集"),
        ("x₍c₎", "候选c是否被选择", "0或1"),
        ("H", "问题3时间资源窗口右端", "643")
    ], widths=[2.8, 8.2, 4.2]
)

page_break()

doc.add_heading("5 问题1 时频冲突检测模型", level=1)
doc.add_heading("5.1 周期时间区间展开", level=2)
body("装备i的周期长度为dᵢ+gᵢ，第k次使用开始时刻与占用区间分别为")
formula("τᵢₖ = sᵢ + k(dᵢ + gᵢ),    Tᵢₖ = [τᵢₖ, τᵢₖ + dᵢ),    k = 0,…,nᵢ-1.")
body("例如A001的首次区间为[35,40)，单次时长5、间隔60、使用3次，因此三个区间为[35,40)、[100,105)和[165,170)。")

doc.add_heading("5.2 冲突判据", level=2)
body("两台装备发生冲突，当且仅当频域存在严格交叠，且至少一对周期时间区间存在严格交叠：")
formula("max(aᵢ,aⱼ) < min(bᵢ,bⱼ),")
formula("∃k,l,    max(τᵢₖ,τⱼₗ) < min(τᵢₖ+dᵢ,τⱼₗ+dⱼ).")
body("一旦任意时段组合满足条件，便将无序装备对{i,j}记为一条冲突边。即使两台装备在多个周期发生交叠，结果表仍只记录一行。")

doc.add_heading("5.3 检测算法与复杂度", level=2)
body("算法遍历C(150,2)=11175个无序装备对。先用常数时间检查频段是否相交；仅对频域相交的装备对枚举nᵢnⱼ个时段组合，发现一次交叠后立即停止。最坏复杂度为O(N²K²)，其中K为最大使用次数。独立复核则建立单位网格occupants[(t,f)]，从同一网格的装备集合中生成冲突对并去重。两种方法所得集合完全一致。")

doc.add_heading("5.4 检测结果", level=2)
add_table(
    ["冲突类别", "冲突装备对数", "占比"],
    [("A-A",0,"0.00%"),("A-B",21,"7.07%"),("A-C",66,"22.22%"),("B-B",10,"3.37%"),("B-C",181,"60.94%"),("C-C",19,"6.40%"),("合计",297,"100.00%")],
    widths=[5.2,5.2,4.2]
)
body("297对冲突占全部11175个装备对的2.66%，但涉及148台装备，占98.67%。仅C007和C060原生无冲突。若按具体周期区间组合计数，共有431次时段交叠；该数不是题目要求的冲突装备对数。B-C冲突181对，是冲突图的主要组成。")
add_figure("图2_冲突类型统计.png", "图3  不同装备类别之间的冲突对统计", width=15.5)
add_figure("图3_装备用频冲突网络.png", "图4  装备用频冲突网络", width=15.7)

doc.add_heading("6 问题2 有限调整下的冲突消解", level=1)
doc.add_heading("6.1 决策变量与动作互斥", level=2)
body("对每台装备设置整数频移uᵢ∈[-10,10]、整数时间平移vᵢ∈[-5,5]、撤销变量zᵢ及调整变量cᵢ。频移与时移至多一个非零；撤销时所有调整量为0。频域需满足0≤aᵢ+uᵢ<bᵢ+uᵢ≤100，首次时间需满足sᵢ+vᵢ≥0。频宽、单次时长、间隔和使用次数保持不变。")

doc.add_heading("6.2 相对位移安全集", level=2)
body("两台装备的周期时间关系只取决于相对时间平移δ=vᵢ-vⱼ。由于δ仅有21种取值，可离线枚举不会使任何周期时段相交的安全集合Sᵢⱼ。任意两台未撤销装备必须满足频域分离或相对时移安全：")
formula("zᵢ=1  ∨  zⱼ=1  ∨  bᵢ+uᵢ≤aⱼ+uⱼ  ∨  bⱼ+uⱼ≤aᵢ+uᵢ  ∨  (vᵢ-vⱼ)∈Sᵢⱼ.")
body("该逻辑是无冲突的充分必要条件。紧凑模型只需建立2142组可能发生接触的装备对约束，避免在求解器内部展开大量单位网格。另建有限候选网格模型作为对照：全网4589个动作候选、35406条资源互斥约束。两种模型共享相同动作边界和冲突定义。")

doc.add_heading("6.3 七层字典序目标", level=2)
body("题面没有给出各目标的数值权重。为使高层目标不被低层收益抵消，按以下顺序逐层求解，每一阶段固定前序目标当前值：")
bullets([
    "最小化撤销总数 Σzᵢ；",
    "最小化A类撤销数；",
    "最小化B类撤销数；",
    "最小化调整总数 Σcᵢ；",
    "最小化A类调整数；",
    "最小化B类调整数；",
    "最小化归一化平移成本 Σ(|uᵢ|+2|vᵢ|)。"
], numbered=True)
body("最后一层相当于10Σ(|uᵢ|/10+|vᵢ|/5)，使频移与时移按各自允许最大幅度归一化。字典序是对题意“尽可能”的一种明确解释，而不是题目唯一指定的偏好。")

doc.add_heading("6.4 结果与最优性状态", level=2)
add_table(
    ["类别", "原样保留", "调整", "撤销", "最终执行"],
    [("A类",4,16,0,20),("B类",2,34,4,36),("C类",12,76,2,88),("合计",18,126,6,144)],
    widths=[3.0,3.0,3.0,3.0,3.0]
)
body("调整中频移91项、时间平移35项；撤销B009、B018、B024、B028、C027、C054。最终执行144项，独立复核残余冲突为0。")
add_table(
    ["目标", "当前值", "下界", "状态"],
    [("撤销总数",6,6,"OPTIMAL"),("A类撤销",0,0,"OPTIMAL"),("B类撤销",4,1,"FEASIBLE"),("调整总数",126,55,"FEASIBLE"),("A类调整",16,8,"FEASIBLE"),("B类调整",34,20,"FEASIBLE"),("平移成本",775,12,"FEASIBLE")],
    widths=[6.0,3.0,3.0,3.5]
)
body("撤销总数6已证明全局最优，说明在题设动作范围内最多执行144项计划；A类撤销0也已证明最优。其余阶段在时限内仅得到可行解，因此不能将完整七层目标向量称为全局字典序最优。")
add_figure("图5_问题2消解结果.png", "图5  问题2各类别保留调整与撤销数量", width=15.5)
add_figure("图7_问题2调整方式与幅度.png", "图6  问题2调整方式和幅度分布", width=15.5)

doc.add_heading("7 问题3 固定方案下的新增容量", level=1)
doc.add_heading("7.1 资源窗口解释", level=2)
body("题目要求不增加时频资源，但未单独给出时间上界。本文以附件1全部原计划的最晚一次结束时刻定义H=643，并固定问题2中144项有效计划。新增C类装备仍使用频宽3、单次时长2、间隔8、使用12次，因此若首次开始为s，其最后一次结束为s+112，合法起点满足s∈{0,…,531}；频率起点f∈{0,…,97}。")

doc.add_heading("7.2 完整候选与集合打包模型", level=2)
body("完整枚举每个(f,s)位置，删除与问题2固定计划冲突的候选，剩余2314个。对每个候选c设置二元变量x₍c₎。对任一单位时频资源格r，记占用该格的候选集合为C(r)，建立")
formula("max  Σ₍c₎ x₍c₎,    s.t.  Σ₍c∈C(r)₎ x₍c₎ ≤ 1,    x₍c₎∈{0,1}.")
body("模型共有7217条去重资源排斥约束。候选枚举覆盖全部整数频率与首次时间起点，单位网格排斥与区间无冲突等价，因此CP-SAT给出的上界可以用于证明新增数量的最优性。")

doc.add_heading("7.3 结果", level=2)
body("最优新增数量为142台，求解器目标值与上界均为142，状态为OPTIMAL。独立验证检查原有与新增、新增与新增之间的全部周期区间，冲突数为0；同时确认所有新增计划落在[0,100)×[0,643)内。该结论依赖于固定的问题2方案和H=643窗口，若改变问题2布局或时间窗口，需重新求解。")
add_figure("图8_问题3新增C类装备布局.png", "图7  固定问题2方案后新增142台C类装备的布局", width=15.7)

doc.add_heading("8 问题4 允许C类间隔调整的冲突消解", level=1)
doc.add_heading("8.1 新增动作与约束", level=2)
body("在问题2动作基础上，C类装备可选择整数间隔变化hᵢ∈[-10,10]，并满足gᵢ+hᵢ≥0。每台装备仍只能选择原样、频移、首次时间平移、间隔调整或撤销中的一种。A、B类间隔不变。调整后的第k次使用区间为")
formula("[sᵢ+vᵢ+k(dᵢ+gᵢ+hᵢ),  eᵢ+vᵢ+k(dᵢ+gᵢ+hᵢ)),    k=0,…,nᵢ-1.")
body("间隔变化会累积改变后续使用时刻，因此必须展开全部周期。幅度成本扩展为Σ(|uᵢ|+2|vᵢ|+|hᵢ|)，仍对应按最大允许幅度归一化后的十倍。")

doc.add_heading("8.2 候选模型与紧凑对照模型", level=2)
body("有限候选模型为每台装备枚举全部合法动作，共6209个候选。每台恰选一个候选；撤销候选不占资源；每个单位时频格最多被一个候选占用，形成42072条去重资源约束。紧凑对照模型把时间平移与间隔变化表示为有限时间模式，预计算装备对之间的安全模式组合，再与频段分离和撤销逻辑合并。两种模型遵守同一动作集合，候选结果按相同七层目标比较。")

doc.add_heading("8.3 结果与解释", level=2)
add_table(
    ["类别", "原样保留", "调整", "撤销", "最终执行"],
    [("A类",4,16,0,20),("B类",2,37,1,39),("C类",5,82,3,87),("合计",11,135,4,146)],
    widths=[3.0,3.0,3.0,3.0,3.0]
)
body("调整包括频移56项、首次时间平移31项、间隔调整48项；撤销B019、C031、C047、C077。方案最终执行146项，独立检测残余冲突为0。与问题2相比，间隔自由度使撤销数量从6降到4，执行数量从144增至146。")
add_table(
    ["目标", "当前值", "下界", "状态"],
    [("撤销总数",4,2,"FEASIBLE"),("A类撤销",0,0,"OPTIMAL"),("B类撤销",1,0,"FEASIBLE"),("调整总数",135,7,"FEASIBLE"),("A类调整",16,6,"FEASIBLE"),("B类调整",37,32,"FEASIBLE"),("归一化成本",826,290,"FEASIBLE")],
    widths=[6.0,3.0,3.0,3.5]
)
body("当前撤销4项尚未证明最优，已知最强下界为2。后续目标也多为限时可行结果。论文只能声称“获得4项撤销的无冲突方案”，不能声称“最少撤销4项”。")
add_figure("图9_问题2与问题4方案对比.png", "图8  问题2与问题4方案对比", width=15.6)
add_figure("图10_问题4调整方式构成.png", "图9  问题4调整方式构成", width=15.4)

doc.add_heading("8.4 时间窗口适用性", level=2)
body("问题4按题面明确条件只限制首次时间非负，没有另加统一结束上限。当前方案中C066、C082、C083、C087的最后结束时刻超过650；因此该方案不能直接解释为统一[0,643)或[0,650)窗口下的答案。若评审要求四问统一时间边界，应在求解器和验证器中同步增加末次结束约束并重新计算问题4；问题3若改用H=650也需重算。")

doc.add_heading("9 模型检验与可靠性分析", level=1)
doc.add_heading("9.1 独立验证闭环", level=2)
body("验证程序不复用求解器的候选生成与资源约束逻辑，而是重新读取原始附件并逐台重建计划。核验内容包括整数性、单参数互斥、撤销与调整互斥、频移和时移幅度、频段边界、首次时间非负、频宽与单次时长不变、间隔权限、使用次数以及全部周期之间的冲突。最后重新读取result1.xlsx至result4.xlsx，检查表头、编号、列位置、空白字段与计算结果的一致性。")

doc.add_heading("9.2 交叉模型验证", level=2)
body("问题1使用区间求交与单位网格两种独立方法；问题2使用相对位移紧凑模型与候选网格模型；问题4使用网格候选模型与时间模式紧凑模型。不同表达得到相同可行性判定，可降低单一建模实现遗漏边界条件的风险。")

doc.add_heading("9.3 最优性与界", level=2)
body("最优性状态是结果解释的一部分。问题2的撤销总数6、A类撤销0以及问题3新增142台已证明最优；问题2后续目标和问题4大多数目标只得到限时可行解。下界与当前值之间的差距反映仍可能存在改进空间，但不影响现有方案的可执行性和无冲突性。")
add_figure("图12_求解结果与最优性状态.png", "图10  各问求解结果与最优性状态", width=15.8)

doc.add_heading("9.4 边界与敏感性讨论", level=2)
add_table(
    ["因素", "当前设定", "对结论的影响"],
    [
        ("区间端点", "左闭右开", "端点相接合法；改为闭区间会增加冲突"),
        ("问题3时间窗口", "H=643", "改变H会改变候选集和最大新增量"),
        ("问题2目标顺序", "七层字典序", "前层严格支配后层；改变顺序可能得到不同方案"),
        ("问题4结束上限", "不额外设置", "若统一限定H，必须重算"),
        ("求解时限与线程", "限时多线程", "可行方案可能变化；已证明最优的目标不受影响")
    ], widths=[3.4,4.2,7.2]
)

doc.add_heading("10 模型评价", level=1)
doc.add_heading("10.1 优点", level=2)
bullets([
    "冲突判据与整数网格完全对应，避免连续区间和离散资源之间的歧义。",
    "相对位移安全集把多周期关系压缩为有限集合，模型规模更紧凑。",
    "字典序目标直接体现少撤销、保护高优先级、少调整和小幅调整。",
    "求解、验证与表格导出分离，结果可复核、可追踪。",
    "对OPTIMAL与FEASIBLE作严格区分，结论强度与证据一致。"
])
doc.add_heading("10.2 局限", level=2)
bullets([
    "问题3的H=643是必要的建模解释，不是题面直接给出的数值。",
    "问题2的后续目标和问题4尚未完成全局最优性证明。",
    "模型未考虑调整实施成本、设备间保护带、随机延迟和需求不确定性。",
    "多目标顺序虽符合题意，但并非唯一合理偏好，需要在实际应用中由管理者确认。"
])
doc.add_heading("10.3 改进方向", level=2)
body("可继续延长主目标证明时间，利用冲突图连通分量分解、对称性破除、拉格朗日下界或列生成强化下界；还可把保护带、设备重要度、调整实施成本和鲁棒时间裕量纳入模型。若必须统一规划窗口，应先确定H，再对问题2至问题4采用一致的末次结束约束。")

doc.add_heading("11 结论", level=1)
body("本文建立了周期用频计划的统一时频建模框架。问题1检测出297对冲突，B-C冲突占主导；问题2在有限频移和时移下获得执行144项的无冲突方案，并证明最少撤销6项；问题3在固定方案和[0,643)时间窗口下最多新增142台C类装备，已证明最优；问题4引入C类间隔调整后获得执行146项、撤销4项的无冲突方案，但最优性尚未证明。独立验证和交叉模型复核保证了约束、结果与提交表的一致性。")

doc.add_heading("参考文献", level=1)
refs = [
    "[1] 全国大学生数学建模竞赛组委会. 2026年高教社杯全国大学生数学建模竞赛D题 时频冲突检测与消解.",
    "[2] Rossi F, van Beek P, Walsh T. Handbook of Constraint Programming. Elsevier, 2006.",
    "[3] Google. OR-Tools CP-SAT Solver Documentation. https://developers.google.com/optimization/cp/cp_solver.",
    "[4] Schrijver A. Theory of Linear and Integer Programming. Wiley, 1998."
]
for x in refs:
    body(x, indent=False)

page_break()

doc.add_heading("附录A 项目思路总览", level=1)
body("本附录供团队复盘、汇报和继续优化使用。最短叙述主线是：先把周期计划展开成时频矩形并建立冲突图；再把允许动作离散化，用CP-SAT在无冲突硬约束下做字典序优化；问题3把剩余空间转化为C类候选集合打包；问题4把C类间隔变化加入时间模式；所有结果均由独立程序重建周期并复核。")
add_table(
    ["问题", "核心输入", "核心模型", "关键结果", "结论强度"],
    [
        ("问题1", "150项原计划", "区间求交+冲突图", "297对冲突", "双方法一致"),
        ("问题2", "原计划+有限平移", "安全相对位移+字典序CP-SAT", "撤销6，执行144", "撤销数已证最优"),
        ("问题3", "固定q2+H=643", "完整候选集合打包", "新增142", "已证最优"),
        ("问题4", "原计划+C类间隔变化", "候选网格+时间模式", "撤销4，执行146", "限时可行，下界2")
    ], widths=[2.0,3.6,4.8,3.0,3.0], font_size=8.5
)
doc.add_heading("A.1 一句话亮点", level=2)
body("用相对时间平移安全集压缩周期冲突关系，以严格字典序处理多重优先级，再用独立区间验证器闭环核验。")
doc.add_heading("A.2 论文必须守住的表述边界", level=2)
bullets([
    "可以说：问题2最少撤销6项已证明；不能说：问题2整套方案七层目标均全局最优。",
    "可以说：问题3在固定q2方案及H=643下新增142台已证明最优；不能脱离该窗口和基准推广。",
    "可以说：问题4找到撤销4项的无冲突方案；不能说4是最少撤销，因为下界仍为2。",
    "问题4没有统一结束上限；若答辩老师要求[0,643)或[0,650)，应承认需补约束重算。"
])
doc.add_heading("A.3 提交前核对清单", level=2)
bullets([
    "摘要中的每个数字与最终Excel和报告一致。",
    "正文不出现队员、学校、赛区等身份信息。",
    "result1.xlsx至result4.xlsx与正文统计一致并通过独立验证。",
    "附录列出支撑材料与必要源程序，排除虚拟环境和node_modules。",
    "如实完成AI工具使用声明和使用详情材料，人工理解并核验关键模型与结果。",
    "正文页数、文件大小、图表清晰度与比赛格式规范一致。"
])

page_break()

doc.add_heading("附录B 答辩高频问答", level=1)
qa = [
    ("1 为什么使用左闭右开区间", "离散资源中相邻计划可在同一端点交接。左闭右开使[1,3)与[3,5)不重叠，也与单位网格占用一致，避免把端点接触误判为冲突。"),
    ("2 间隔时长如何理解", "间隔是前一次使用结束到下一次使用开始的空闲时长，因此相邻两次开始时刻之差为单次时长d加间隔g，而不是仅为g。"),
    ("3 为什么297对而不是431次", "297是不同装备对的数量，符合result1模板；431是具体周期时段组合的相交次数。同一装备对可在多个周期冲突，但只输出一次。"),
    ("4 为什么先判断频率再判断时间", "冲突必须同时满足两维交叠。频段检查是常数时间，能提前剔除大量装备对，减少周期时段枚举。"),
    ("5 为什么原冲突图不能直接作为问题2全部约束", "调整可能让原本不冲突的装备产生新冲突，所以优化模型必须检查所有潜在装备对，而不是只消除原297条边。"),
    ("6 为什么选择CP-SAT而不是遗传算法", "变量和候选均为离散整数，冲突是逻辑互斥，CP-SAT与问题结构匹配，并能给出上下界和最优性状态。遗传算法可找可行解，但通常不能证明最少撤销数。"),
    ("7 相对位移安全集Sᵢⱼ是什么", "它列出两台装备在所有周期上均不发生时间交叠的相对时移δ=vᵢ-vⱼ。时移范围有限，所以只需预计算21种δ。"),
    ("8 相对位移模型为什么是精确的", "平移不会改变单次时长、间隔和次数，两台计划的全部时间关系只由相对平移决定。枚举全部δ并逐周期检验后得到的安全集既不漏解也不加入伪解。"),
    ("9 单参数调整如何保证", "每台装备的动作在原样、频移、时移和撤销候选中恰选一个；紧凑模型中则通过非零指示和撤销逻辑保证u与v不同时变化。"),
    ("10 为什么采用字典序而不是加权和", "题面强调先少撤销、保护高优先级、再少调整和小幅调整，却没有给权重。字典序能保证低层收益永远不会抵消高层目标，更符合定性优先关系。"),
    ("11 七层顺序是不是题目唯一答案", "不是。它是对定性要求的明确建模选择。我们在论文中公开顺序，并用同一顺序比较问题2和问题4；若管理偏好改变，可以调整层次后重算。"),
    ("12 为什么时间平移成本乘2", "允许频移最大10格、时移最大5格。|u|+2|v|等于归一化成本10(|u|/10+|v|/5)，用统一无量纲尺度比较两类幅度。"),
    ("13 问题2哪些结果真正最优", "撤销总数6和A类撤销0已证明最优。B类撤销、调整数量和幅度阶段只取得限时可行解，不能称完整方案全局字典序最优。"),
    ("14 如何证明问题2至少撤销6项", "CP-SAT在同一完整可行域上求解第一层目标，得到可行值6且下界也为6，目标值与下界闭合，因此任何合法方案都不能少于6项撤销。"),
    ("15 为什么问题3需要规定时间窗口", "如果时间轴没有上界，新装备可以不断向后平移，最大数量可能失去有限意义。本文把不增加资源解释为沿用原计划覆盖窗口[0,643)。"),
    ("16 为什么H取643", "643是附件1中所有原始计划展开全部使用次数后的最晚结束时刻。它由数据计算得到，不是题面单独给出的常数，因此结论需附带这一假设。"),
    ("17 为什么问题3新增计划可以同时选频率和时间", "新增装备没有原始坐标，不存在从原计划只能改一个参数的问题。模型是在既定资源窗口内为新计划选择一个频段起点和首次时间起点。"),
    ("18 142台为什么是最大值", "模型完整枚举全部整数合法起点，删除与固定计划冲突的候选，再对候选间资源互斥做精确0-1优化。求解器目标值和上界均为142，因此在当前假设下已证明最优。"),
    ("19 为什么不能用资源面积上界证明142", "面积上界忽略频段连续性、12次周期结构和固定计划造成的碎片，数值很松。142的证明来自精确候选模型的上界，不来自面积估计。"),
    ("20 问题4间隔变化为何能少撤销", "间隔变化会改变后续各次使用的位置，提供比整体时移更丰富的时间模式，使部分原本无法通过±5整体平移消解的冲突可以错开。"),
    ("21 间隔变化会不会与频移或时移同时发生", "不会。候选集合要求每台只选一种操作；C类可以单独改间隔，但不能在同一方案行同时频移、时移或改间隔。"),
    ("22 问题4的4项撤销是不是最优", "不是已证最优。当前方案撤销4项，已知下界为2，状态FEASIBLE。准确表述是找到了4项撤销的无冲突方案。"),
    ("23 为什么问题4不沿用H=643", "H=643是为问题3“不增加资源”作出的专门解释。问题4题面只增加间隔调整权限，并未明确统一结束上限。若评审要求统一窗口，应增加末次结束约束后重算。"),
    ("24 如何保证Excel没有填错", "验证器重新读取四个结果表，核对模板列、编号、撤销标记和空白字段，并把表中动作还原为完整计划，再与独立冲突检测结果比较。"),
    ("25 独立验证为何可信", "验证程序从原始附件重新构造周期区间，不调用求解模型的候选与冲突约束函数；因此能发现求解建模、导出或字段映射中的独立错误。"),
    ("26 模型最大的创新点是什么", "一是用安全相对时移集合压缩多周期冲突关系；二是用严格字典序解释多重优先级；三是把求解与独立验证分离，形成可复现闭环。"),
    ("27 模型最大的局限是什么", "问题3窗口需要解释，问题4最优性尚未闭合；同时尚未考虑保护带、设备切换成本、随机延迟和不确定需求。"),
    ("28 如果继续优化最值得做什么", "优先加强问题4撤销数的下界或找到更少撤销方案；其次在固定前层目标下减少B类撤销与调整总数；若统一时间窗口，则先重算边界一致的q3和q4。"),
    ("29 结果如何复现", "从附件1读取输入，依次运行q1检测、q2求解与验证、固定q2后的q3求解、q4重新求解，再导出四个模板表。每问均保留求解状态、界和独立验证记录。"),
    ("30 答辩时最容易说错什么", "把FEASIBLE说成最优、把431说成冲突对数、忽略间隔定义中的单次时长、忘记问题3的H=643假设，或把问题4方案误称为统一时间窗口内可行。")
]
for q, a in qa:
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.first_line_indent = Cm(0)
    r = p.add_run(q)
    set_font(r, east="黑体", size=10.5, bold=True)
    p2 = doc.add_paragraph()
    p2.paragraph_format.first_line_indent = Cm(0.74)
    p2.paragraph_format.line_spacing = 1.4
    p2.paragraph_format.space_after = Pt(3)
    r = p2.add_run("答：" + a)
    set_font(r, size=10.2)

doc.add_heading("附录C 支撑材料说明", level=1)
body("最终支撑材料应包含四问源程序、依赖声明、最终result1.xlsx至result4.xlsx、必要的求解状态和验证记录。虚拟环境、node_modules及临时缓存不属于必要源程序，不宜打包。论文正文与支撑材料应分别检查匿名性、大小限制和可复现性。")

# 统一段落孤行控制和表格字体
for p in doc.paragraphs:
    p.paragraph_format.widow_control = True

# 文档元数据
props = doc.core_properties
props.title = "离散时频资源下用频计划冲突检测与多目标消解"
props.subject = "D题论文初稿及答辩问答"
props.keywords = "时频冲突, CP-SAT, 字典序优化, 集合打包"
props.author = ""
props.last_modified_by = ""

doc.save(OUTPUT)
print(OUTPUT)
