"""Independent verifier for the question 2 solution artifact.

This module deliberately does not import the question 2 solver.  It rebuilds
the source plans, applies the operations recorded in solution.json, and checks
the resulting schedule by pairwise interval enumeration.
"""
from __future__ import annotations

import json
import math
import re
from collections import Counter
from pathlib import Path

import openpyxl


ROOT = Path(__file__).resolve().parents[3]
SOURCE = ROOT / "附件" / "附件1.xlsx"
SOLUTION = ROOT / "answer" / ".cache" / "q2" / "solution.json"
REPORT = ROOT / "answer" / ".cache" / "q2" / "verification.json"
RESULT = ROOT / "answer" / "results" / "result2.xlsx"
HEADERS = ["用频装备编号", "调整后频段区间", "调整后时间区间", "是否撤销用频计划"]


def interval(value):
    nums = re.findall(r"-?\d+(?:\.\d+)?", str(value))
    if len(nums) != 2:
        raise ValueError(f"invalid interval: {value!r}")
    return tuple(int(float(x)) for x in nums)


def load_source():
    ws = openpyxl.load_workbook(SOURCE, data_only=True).active
    out = {}
    for row in list(ws.values)[1:]:
        ident, freq, time, gap, count = row
        lo, hi = interval(freq)
        start, end = interval(time)
        out[ident] = dict(id=ident, lo=lo, hi=hi, start=start, end=end,
                          gap=gap, count=count)
    return out


def num(value, default=0):
    if value is None or value == "":
        return default
    return float(value)


def same(a, b):
    return a == b or (isinstance(a, (int, float)) and isinstance(b, (int, float))
                      and math.isclose(a, b, abs_tol=1e-9))


def fmt_interval(a, b):
    return f"[{int(a) if a == int(a) else a},{int(b) if b == int(b) else b})"


def row_values(row):
    if isinstance(row, dict):
        return [row.get(k) for k in ("id", "freq", "time", "canceled")]
    return list(row) if isinstance(row, (list, tuple)) else [row]


def main():
    report = {"source": str(SOURCE), "solution": str(SOLUTION), "checks": {}, "errors": []}
    source = load_source()
    report["source_count"] = len(source)
    if not SOLUTION.exists():
        report["errors"].append("solution.json does not exist")
        report["checks"]["solution_present"] = False
        REPORT.parent.mkdir(parents=True, exist_ok=True)
        REPORT.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
        return 1
    report["checks"]["solution_present"] = True
    try:
        solution = json.loads(SOLUTION.read_text(encoding="utf-8"))
    except Exception as exc:
        report["errors"].append(f"cannot parse solution.json: {exc}")
        REPORT.parent.mkdir(parents=True, exist_ok=True)
        REPORT.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
        return 1

    plans = solution.get("plans", [])
    by_id = {p.get("id"): p for p in plans if isinstance(p, dict)}
    errors = report["errors"]
    if len(plans) != len(source) or len(by_id) != len(plans) or set(by_id) != set(source):
        errors.append("plans must contain exactly one record for every source equipment")
    report["checks"]["plans_complete"] = not errors

    operations = {}
    for ident, base in source.items():
        p = by_id.get(ident)
        if p is None:
            continue
        canceled = p.get("canceled", False)
        if not isinstance(canceled, bool):
            errors.append(f"{ident}: canceled is not boolean")
            canceled = bool(canceled)
        df, dt = num(p.get("df")), num(p.get("dt"))
        if not math.isfinite(df) or not math.isfinite(dt) or not df.is_integer() or not dt.is_integer():
            errors.append(f"{ident}: df and dt must be finite integers")
        # The plan record stores the resulting bounds; df/dt store the move.
        for key, expected in (("lo", base["lo"] + df), ("hi", base["hi"] + df),
                              ("start", base["start"] + dt), ("end", base["end"] + dt),
                              ("gap", base["gap"]), ("count", base["count"])):
            if not same(p.get(key), expected):
                errors.append(f"{ident}: resulting field {key} disagrees with source and shift")
        if abs(df) > 10 + 1e-9 or abs(dt) > 5 + 1e-9:
            errors.append(f"{ident}: shift exceeds allowed magnitude")
        if abs(df) > 1e-9 and abs(dt) > 1e-9:
            errors.append(f"{ident}: frequency and time shifts are simultaneous")
        if canceled and (abs(df) > 1e-9 or abs(dt) > 1e-9):
            errors.append(f"{ident}: canceled plan also has a shift")
        operations[ident] = (p, canceled, df, dt)

    # Validate resulting plans and independently enumerate all retained conflicts.
    retained = []
    for ident, (p, canceled, df, dt) in operations.items():
        lo, hi = p["lo"], p["hi"]
        start, end = p["start"], p["end"]
        if lo < -1e-9 or hi > 100 + 1e-9 or lo >= hi:
            errors.append(f"{ident}: adjusted frequency is outside [0,100]")
        if start < -1e-9 or end < start:
            errors.append(f"{ident}: adjusted first interval is invalid/nonnegative")
        if canceled:
            continue
        retained.append((ident, lo, hi, start, end, p["gap"], p["count"]))
    conflicts = []
    for i, a in enumerate(retained):
        for b in retained[i + 1:]:
            if max(a[1], b[1]) >= min(a[2], b[2]):
                continue
            da, db = a[4] - a[3], b[4] - b[3]
            for k in range(int(a[6])):
                for j in range(int(b[6])):
                    sa, sb = a[3] + k * (da + a[5]), b[3] + j * (db + b[5])
                    if max(sa, sb) < min(sa + da, sb + db):
                        conflicts.append((a[0], b[0]))
                        break
                else:
                    continue
                break
    report["checks"]["retained_plans_zero_conflict"] = not conflicts
    if conflicts:
        errors.append(f"retained plans have {len(conflicts)} conflicting equipment pairs")

    derived = Counter()
    for _, canceled, df, dt in operations.values():
        if canceled:
            derived["canceled"] += 1
        elif abs(df) > 1e-9:
            derived["frequency"] += 1
            derived["adjusted"] += 1
        elif abs(dt) > 1e-9:
            derived["time"] += 1
            derived["adjusted"] += 1
        else:
            derived["unchanged"] += 1
    stats = solution.get("stats", {})
    for cat in "ABC":
        cat_derived = Counter()
        for ident, (_, canceled, df, dt) in operations.items():
            if ident[0] != cat:
                continue
            if canceled:
                cat_derived["canceled"] += 1
            elif abs(df) > 1e-9:
                cat_derived["frequency"] += 1; cat_derived["adjusted"] += 1
            elif abs(dt) > 1e-9:
                cat_derived["time"] += 1; cat_derived["adjusted"] += 1
            else:
                cat_derived["unchanged"] += 1
        for key in ("unchanged", "adjusted", "canceled", "frequency", "time"):
            expected = cat_derived[key]
            got = stats.get(cat, {}).get(key) if isinstance(stats.get(cat), dict) else None
            if got != expected:
                errors.append(f"stats.{cat}.{key}: expected {expected}, got {got}")
    report["derived_stats"] = dict(derived)
    report["checks"]["stats_match"] = not any(x.startswith("stats.") for x in errors)

    # Rows contain only changed or revoked plans; D is blank for retained plans.
    expected_rows = []
    for ident, (p, canceled, df, dt) in operations.items():
        if canceled or abs(df) > 1e-9 or abs(dt) > 1e-9:
            expected_rows.append([ident, None if canceled or abs(df) <= 1e-9 else fmt_interval(p["lo"], p["hi"]),
                                  None if canceled or abs(dt) <= 1e-9 else fmt_interval(p["start"], p["end"]),
                                  "是" if canceled else None])
    got_rows = solution.get("rows")
    canonical = [row_values(x) for x in got_rows] if isinstance(got_rows, list) else None
    report["checks"]["solution_rows_match"] = canonical == expected_rows
    if canonical != expected_rows:
        errors.append("solution.rows does not exactly match changed/revoked plans")

    if RESULT.exists():
        book = openpyxl.load_workbook(RESULT, data_only=True)
        ws = book.active
        actual = [[ws.cell(r, c).value for c in range(1, 5)] for r in range(1, ws.max_row + 1)]
        ok = (book.sheetnames == ["Sheet1"] and ws.max_column == 4 and
              actual[:1] == [HEADERS] and actual[1:] == expected_rows)
        report["checks"]["excel_match"] = ok
        if not ok:
            errors.append("result2.xlsx headers or rows do not exactly match expected output")
    else:
        report["checks"]["excel_match"] = None

    report["ok"] = not errors
    REPORT.parent.mkdir(parents=True, exist_ok=True)
    REPORT.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0 if report["ok"] else 1


if __name__ == "__main__":
    import argparse

    parser = argparse.ArgumentParser(description="Independently validate a q2 solution artifact")
    parser.add_argument("--solution", default=None, help="solution filename inside .cache/q2")
    parser.add_argument("--report", default=None, help="verification filename inside .cache/q2")
    parser.add_argument("--skip-excel", action="store_true",
                        help="validate the JSON plan only; do not compare the submission workbook")
    args = parser.parse_args()
    if args.solution:
        SOLUTION = SOLUTION.parent / args.solution
    if args.report:
        REPORT = REPORT.parent / args.report
    if args.skip_excel:
        RESULT = None
    raise SystemExit(main())
