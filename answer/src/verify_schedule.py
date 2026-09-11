"""Independent verification from source XLSX and exported submission files.

Does not import solver code or its grid/candidate helpers. Conflicts are checked
by enumerating pairs of half-open intervals, not by the solver's resource grid.
"""
import hashlib
import json
import re
from collections import Counter
from itertools import combinations
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
ANSWER = ROOT / 'answer'


def require(condition, message):
    if not condition:
        raise ValueError(message)


def parse_interval(value):
    match = re.fullmatch(r'\[(-?\d+),(-?\d+)\)', str(value).replace(' ', ''))
    require(match is not None, f'Invalid integer half-open interval: {value!r}')
    return tuple(map(int, match.groups()))


def source_plans():
    book = openpyxl.load_workbook(ROOT / '附件/附件1.xlsx', data_only=True)
    out = {}
    for ident, freq, time, gap, count in list(book.active.values)[1:]:
        lo, hi = parse_interval(freq)
        start, end = parse_interval(time)
        require(ident not in out, f'Duplicate source ID: {ident}')
        out[ident] = dict(id=ident, lo=lo, hi=hi, start=start, end=end, gap=gap, count=count)
    book.close()
    return out


def periods(p):
    duration = p['end'] - p['start']
    return [(p['start'] + k * (duration + p['gap']),
             p['end'] + k * (duration + p['gap'])) for k in range(p['count'])]


def check_conflicts(plans):
    expanded = [(p, periods(p)) for p in plans if not p.get('canceled', False)]
    for (a, ta), (b, tb) in combinations(expanded, 2):
        if max(a['lo'], b['lo']) >= min(a['hi'], b['hi']):
            continue
        require(not any(max(sa, sb) < min(ea, eb) for sa, ea in ta for sb, eb in tb),
                f"Conflict: {a['id']} / {b['id']}")
    return len(expanded) * (len(expanded) - 1) // 2


def check_operations(source, plans, question):
    require(len(plans) == len(source) and {p['id'] for p in plans} == set(source),
            'Plans must contain every original equipment exactly once')
    stats = {g: Counter() for g in 'ABC'}
    rows = []
    for p in plans:
        base = source[p['id']]
        df, dt, dg = p['df'], p['dt'], p.get('dg', 0)
        require(all(type(v) is int for v in (df, dt, dg)), f"{p['id']}: noninteger shift")
        require(type(p['canceled']) is bool, f"{p['id']}: canceled must be boolean")
        require(abs(df) <= 10 and abs(dt) <= 5 and abs(dg) <= 10, f"{p['id']}: shift limit")
        require(sum(v != 0 for v in (df, dt, dg)) <= 1, f"{p['id']}: multiple parameters changed")
        require(not dg or (question == 4 and p['id'].startswith('C')), f"{p['id']}: forbidden gap adjustment")
        require(not p['canceled'] or (df, dt, dg) == (0, 0, 0), f"{p['id']}: canceled and changed")
        expected = dict(base, lo=base['lo'] + df, hi=base['hi'] + df,
                        start=base['start'] + dt, end=base['end'] + dt, gap=base['gap'] + dg)
        require(all(type(p[k]) is int and p[k] == expected[k]
                    for k in ('lo', 'hi', 'start', 'end', 'gap', 'count')),
                f"{p['id']}: final parameters disagree with original plan and operation")
        require(0 <= p['lo'] < p['hi'] <= 100 and 0 <= p['start'] < p['end'] and p['gap'] >= 0,
                f"{p['id']}: invalid resource bounds")
        s = stats[p['id'][0]]
        if p['canceled']:
            s['canceled'] += 1
            rows.append([p['id'], None, None, None, '是'])
        elif df or dt or dg:
            s['adjusted'] += 1
            s['frequency' if df else 'time' if dt else 'gap'] += 1
            rows.append([p['id'], f"[{p['lo']},{p['hi']})" if df else None,
                         f"[{p['start']},{p['end']})" if dt else None,
                         p['gap'] if dg else None, None])
        else:
            s['unchanged'] += 1
    return rows, stats


def check_excel(question, rows):
    template = openpyxl.load_workbook(ROOT / f'附件/附件2/result{question}.xlsx')
    output_path = ANSWER / f'results/result{question}.xlsx'
    require(output_path.exists(), f'Missing submission: {output_path}')
    book = openpyxl.load_workbook(output_path, data_only=False)
    require(book.sheetnames == template.sheetnames, 'Worksheet names changed')
    expected_header = list(next(template.active.values))
    actual = list(map(list, book.active.values))
    require(actual == [expected_header] + rows, f'result{question}.xlsx headers/rows/blank fields mismatch')
    require(book.active.max_column == len(expected_header), 'Unexpected result columns')
    template.close()
    book.close()


def check_q3(data, source):
    q2_path = ANSWER / '.cache/q2/solution.json'
    require(data['q2_sha256'] == hashlib.sha256(q2_path.read_bytes()).hexdigest(),
            'Q2 baseline changed; rerun Q3 before export')
    require(data['source_sha256'] == hashlib.sha256((ROOT / '附件/附件1.xlsx').read_bytes()).hexdigest(),
            'Source changed; rerun Q3')
    q2 = json.loads(q2_path.read_text(encoding='utf-8'))
    q2_rows, _ = check_operations(source, q2['plans'], 2)
    check_excel(2, [r[:3] + r[4:] for r in q2_rows])
    fixed = [p for p in q2['plans'] if not p['canceled']]
    horizon = max(end for p in source.values() for _, end in periods(p))
    require(data['horizon'] == horizon, 'Incorrect resource horizon')
    shapes = {(p['hi'] - p['lo'], p['end'] - p['start'], p['gap'], p['count'])
              for p in source.values() if p['id'].startswith('C')}
    require(len(shapes) == 1, 'Ambiguous C profile')
    shape = shapes.pop()
    require(data['profile'] == dict(zip(('width', 'duration', 'gap', 'count'), shape)), 'Wrong C profile metadata')
    plans = data['plans']
    require(len(plans) == data['count'], 'Wrong added equipment count')
    require([p['id'] for p in plans] == list(range(1, len(plans) + 1)), 'Invalid new equipment sequence')
    for p in plans:
        require(all(type(p[k]) is int for k in ('lo', 'hi', 'start', 'end', 'gap', 'count')), 'Noninteger new plan')
        require((p['hi'] - p['lo'], p['end'] - p['start'], p['gap'], p['count']) == shape,
                f"New plan {p['id']}: C requirements changed")
    for p in fixed + plans:
        require(0 <= p['lo'] < p['hi'] <= 100 and p['start'] >= 0 and periods(p)[-1][1] <= horizon,
                f"{p['id']}: resource window exceeded")
    rows = [[p['id'], f"[{p['lo']},{p['hi']})", f"[{p['start']},{p['end']})"] for p in plans]
    require(data['rows'] == rows, 'Q3 cached rows mismatch')
    require(data['status'] in ('OPTIMAL', 'FEASIBLE') and data['count'] <= data['upper_bound'], 'Invalid objective bound')
    if data['status'] == 'OPTIMAL':
        require(data['count'] == data['upper_bound'], 'Optimum and bound disagree')
    check_excel(3, rows)
    pairs = check_conflicts(fixed + plans)
    return dict(fixed=len(fixed), added=len(plans), horizon=horizon, pair_checks=pairs, conflicts=0,
                baseline_unchanged=True, excel_match=True)


def check_q4(data, source):
    rows, stats = check_operations(source, data['plans'], 4)
    require(data['rows'] == rows, 'Q4 cached rows mismatch')
    for group in 'ABC':
        for key in ('unchanged', 'adjusted', 'canceled', 'frequency', 'time', 'gap'):
            require(data['stats'][group][key] == stats[group][key], f'Incorrect {group}.{key} statistic')
    cost = sum(abs(p['df']) + 2 * abs(p['dt']) + abs(p['dg']) for p in data['plans'])
    require(data['shift_cost'] == cost, 'Incorrect shift cost')
    objectives = dict(cancel_total=sum(s['canceled'] for s in stats.values()),
                      cancel_A=stats['A']['canceled'], cancel_B=stats['B']['canceled'],
                      adjust_total=sum(s['adjusted'] for s in stats.values()),
                      adjust_A=stats['A']['adjusted'], adjust_B=stats['B']['adjusted'], shift_cost=cost)
    for stage in data['stages']:
        if 'value' in stage:
            require(stage['value'] == objectives[stage['objective']], 'Stage value does not match final plan')
            require(stage['lower_bound'] <= stage['value'] + 1e-6, 'Invalid minimization bound')
            if stage['status'] == 'OPTIMAL':
                require(abs(stage['lower_bound'] - stage['value']) < 1e-6, 'Optimal stage has nonzero gap')
    check_excel(4, rows)
    pairs = check_conflicts(data['plans'])
    return dict(stats={g: dict(s) for g, s in stats.items()}, pair_checks=pairs, conflicts=0, excel_match=True)


def verify(question):
    cache = ANSWER / f'.cache/q{question}'
    report = dict(ok=False, question=question, errors=[])
    try:
        data = json.loads((cache / 'solution.json').read_text(encoding='utf-8'))
        report.update((check_q3 if question == 3 else check_q4)(data, source_plans()))
        report['solution_sha256'] = hashlib.sha256((cache / 'solution.json').read_bytes()).hexdigest()
        report['result_sha256'] = hashlib.sha256((ANSWER / f'results/result{question}.xlsx').read_bytes()).hexdigest()
        report['ok'] = True
    except (ValueError, KeyError, TypeError, OSError) as exc:
        report['errors'].append(str(exc))
    cache.mkdir(parents=True, exist_ok=True)
    (cache / 'verification.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0 if report['ok'] else 1
