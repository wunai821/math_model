"""使用项目已有 Python 依赖导出四问 Excel 结果。"""
from __future__ import annotations

import argparse
import json
from pathlib import Path

import openpyxl


ROOT = Path(__file__).resolve().parents[2]
ANSWER = ROOT / "answer"


def load_rows(question: int) -> list[list[object]]:
    # 问题1保存的是冲突对，问题2至4保存的是最终提交行。
    cache = ANSWER / ".cache" / f"q{question}"
    if question == 1:
        data = json.loads((cache / "data.json").read_text(encoding="utf-8"))
        return [[index, *pair] for index, pair in enumerate(data["pairs"], start=1)]
    data = json.loads((cache / "solution.json").read_text(encoding="utf-8"))
    return [list(row) for row in data["rows"]]


def export(question: int) -> Path:
    template = ROOT / "附件" / "附件2" / f"result{question}.xlsx"
    output = ANSWER / "results" / f"result{question}.xlsx"
    rows = load_rows(question)

    book = openpyxl.load_workbook(template)
    sheet = book.active
    expected_columns = sheet.max_column
    if any(len(row) != expected_columns for row in rows):
        raise ValueError(
            f"Question {question} has rows with a column count different from the template"
        )

    # 先清理旧数据，避免不同次限时求解的行数不同而残留旧记录。
    if sheet.max_row > 1:
        for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row,
                                   max_col=expected_columns):
            for cell in row:
                cell.value = None

    # 从第二行开始写入结果，第一行模板表头保持不变。
    for row_index, values in enumerate(rows, start=2):
        for column_index, value in enumerate(values, start=1):
            sheet.cell(row=row_index, column=column_index, value=value)

    output.parent.mkdir(parents=True, exist_ok=True)
    book.save(output)
    book.close()
    print(output)
    return output


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("question", type=int, choices=(1, 2, 3, 4))
    args = parser.parse_args()
    export(args.question)
